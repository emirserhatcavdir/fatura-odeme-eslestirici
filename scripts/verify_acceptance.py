r"""Bağımsız kabul dosyalarını gerçek uygulama okuyucu/yazıcılarıyla doğrular.

Çalıştırma: .venv\Scripts\python.exe -m scripts.verify_acceptance
Testler geçici klasör kullanır; bu komut acceptance_data/reports içine rapor yazar.
"""

import argparse
import csv
from datetime import date, datetime
from decimal import Decimal
import json
from pathlib import Path

from openpyxl import load_workbook

from eslestirici.exports import csv_bytes, frame, report_sheets, xlsx_bytes
from eslestirici.reconciliation import INVOICE_RESULT_COLUMNS, PAYMENT_COLUMNS, process_files, summarize
from eslestirici.views import filter_invoices, invoice_detail

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "acceptance_data"

# Excel başlıklarını bağımsız sabit beklenti alanlarına bağlar; CSV şeması aynı kalır.
EXCEL_FIELDS = {
    "Fatura no": "fatura_no", "Müşteri": "musteri", "Fatura tarihi": "fatura_tarihi",
    "Vade tarihi": "vade_tarihi", "Fatura tutarı (TL)": "tutar_tl",
    "Ödenen (TL)": "toplam_odeme_tl", "Kalan borç (TL)": "kalan_borc_tl",
    "Fazla ödeme (TL)": "fazla_odeme_tl", "Durum": "durum",
    "Gecikme günü": "gecikme_gun", "Gecikme grubu": "gecikme_grubu",
    "Ödeme kimliği": "odeme_id", "Ödeme tarihi": "odeme_tarihi",
    "Ödeme tutarı (TL)": "tutar_tl", "Açıklama": "neden",
    "Gösterge": "gosterge", "Tutar (TL)": "tutar_tl",
    "Gecikmiş borç (TL)": "borc_tl", "Alan": "alan", "Değer": "deger",
    "Müşteri kimliği": "musteri_id", "Eşleşen ödeme (TL)": "toplam_odeme_tl",
    "Kalan alacak (TL)": "kalan_borc_tl", "Gecikmiş alacak (TL)": "gecikmis_alacak_tl",
    "Açık fatura sayısı": "acik_fatura_sayisi", "En eski gecikme (gün)": "en_eski_gecikme_gun",
}


def expected_for(report_date: str) -> dict:
    expected = json.loads((DATA / "expected_results.json").read_text(encoding="utf-8"))
    expected["sheet_names"] = ["Özet", "Rapor Bilgisi", "Müşteri Özeti", "Faturalar", "Eşleşen Ödemeler", "Eşleşmeyen Ödemeler", "Gelecek Faturalar", "Gelecek Ödemeler", "Gecikme Dağılımı"]
    if report_date == "2026-09-26":
        # Kullanıcının ikinci tarih için verdiği sabitler; uygulamadan türetilmez.
        expected["report_date"] = "2026-09-26"
        expected["summary_cents"] = {"Toplam fatura": 500000, "Eşleşen ödeme": 385000, "Kalan borç": 125000, "Gecikmiş borç": 125000, "Fazla ödeme": 10000}
        expected["summary_tl"] = {"Toplam fatura": "5000.00", "Eşleşen ödeme": "3850.00", "Kalan borç": "1250.00", "Gecikmiş borç": "1250.00", "Fazla ödeme": "100.00"}
        expected["invoices"][1].update(gecikme_gun=98, gecikme_grubu="90 üzeri gün")
        expected["invoices"][2].update(toplam_odeme_kurus=150000, kalan_borc_kurus=0, durum="Ödendi", toplam_odeme_tl="1500.00", kalan_borc_tl="0.00")
        expected["matched_payment_ids"] = ["P001", "P002", "P003", "P004", "P006"]
        expected["payments_by_invoice"]["0003"] = ["P006"]
        expected["future_payments"] = []
        expected["aging_cents"] = {"1–30 gün": 0, "31–60 gün": 0, "61–90 gün": 0, "90 üzeri gün": 125000}
    elif report_date != "2026-06-30":
        raise ValueError("Kabul senaryosu yalnızca 2026-06-30 ve 2026-09-26 için tanımlı.")
    return expected


def normalized(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def assert_rows(actual: list[dict], expected: list[dict]) -> None:
    assert len(actual) == len(expected), f"Kayıt sayısı farklı: {len(actual)} / {len(expected)}"
    for actual_row, expected_row in zip(actual, expected):
        assert {key: normalized(value) for key, value in actual_row.items()} == {
            key: normalized(value) for key, value in expected_row.items()
        }, f"Beklenmeyen satır: {actual_row}"


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sheet_rows(sheet) -> list[dict]:
    header_row = 3 if sheet.title == "Özet" else 1
    rows = list(sheet.iter_rows(min_row=header_row, values_only=True))
    keys = [EXCEL_FIELDS.get(value, value) for value in rows[0]]
    return [{key: format(Decimal(str(value)), ".2f") if key.endswith("_tl") else value for key, value in zip(keys, row)} for row in rows[1:]]


def assert_excel_cell_types(workbook) -> None:
    for sheet in workbook:
        header_row = 3 if sheet.title == "Özet" else 1
        columns = {EXCEL_FIELDS.get(cell.value, cell.value): cell.column for cell in sheet[header_row]}
        for key, index in columns.items():
            letter = sheet.cell(header_row, index).column_letter
            if key.endswith("_kurus"):
                assert sheet.column_dimensions[letter].hidden
            for row in sheet.iter_rows(min_row=header_row + 1):
                cell = row[index - 1]
                if key.endswith("_tl"):
                    cents = row[columns[key.removesuffix("_tl") + "_kurus"] - 1]
                    assert cell.data_type == cents.data_type == "n"
                    assert Decimal(str(cell.value)) * 100 == cents.value
                    assert cell.number_format == '#,##0.00 "TL"'
                elif key in ("fatura_no", "odeme_id"):
                    assert cell.data_type == "s" and cell.number_format == "@"
                elif key.endswith("_tarihi"):
                    assert cell.data_type == "d" and cell.number_format == "dd.mm.yyyy"


def verify_case(extension: str, output_folder: Path, report_date: str = "2026-06-30") -> dict:
    expected = expected_for(report_date)
    invoice_path, payment_path = DATA / f"faturalar.{extension}", DATA / f"odemeler.{extension}"
    report, errors = process_files(invoice_path.read_bytes(), invoice_path.name, payment_path.read_bytes(), payment_path.name, date.fromisoformat(expected["report_date"]))
    assert not errors, errors
    assert report is not None
    assert summarize(report.faturalar) == expected["summary_cents"]
    assert report.rapor_tarihi.isoformat() == expected["report_date"]
    expected_core = [{key: value for key, value in row.items() if not key.endswith("_tl")} for row in expected["invoices"]]
    assert_rows(report.faturalar, expected_core)
    assert [row["odeme_id"] for row in report.eslesen_odemeler] == expected["matched_payment_ids"]
    for identifier, ids in expected["payments_by_invoice"].items():
        assert [row["odeme_id"] for row in invoice_detail(report, identifier)["odemeler"]] == ids
    assert report.gelecek_faturalar == []

    output_folder.mkdir(parents=True, exist_ok=True)
    invoices = frame(report.faturalar, INVOICE_RESULT_COLUMNS)
    filtered = filter_invoices(invoices, query="Beta")
    assert filtered["fatura_no"].tolist() == expected["filtered_invoice_ids"]
    # Üretimde kullanılan işlevler. Filtreli veri yalnızca CSV'ye gider.
    (output_folder / "fatura_raporu.csv").write_bytes(csv_bytes(invoices))
    (output_folder / "filtreli_faturalar.csv").write_bytes(csv_bytes(filtered))
    (output_folder / "eslesmeyen_odemeler.csv").write_bytes(csv_bytes(frame(report.eslesmeyen_odemeler, PAYMENT_COLUMNS + ["neden"])))
    (output_folder / "tum_rapor.xlsx").write_bytes(xlsx_bytes(report_sheets(report)))

    assert_rows(read_csv(output_folder / "fatura_raporu.csv"), expected["invoices"])
    expected_filtered = [row for row in expected["invoices"] if row["fatura_no"] in expected["filtered_invoice_ids"]]
    assert_rows(read_csv(output_folder / "filtreli_faturalar.csv"), expected_filtered)
    assert_rows(read_csv(output_folder / "eslesmeyen_odemeler.csv"), expected["unmatched"])

    workbook = load_workbook(output_folder / "tum_rapor.xlsx", data_only=False)
    try:
        assert workbook.sheetnames == expected["sheet_names"]
        assert_excel_cell_types(workbook)
        assert normalized(workbook["Özet"]["B1"].value) == expected["report_date"]
        assert workbook["Özet"]["B1"].data_type == "d"
        info = {row["alan"]: row["deger"] for row in sheet_rows(workbook["Rapor Bilgisi"])}
        assert normalized(info["Raporlama tarihi"]) == expected["report_date"]
        summary = sheet_rows(workbook["Özet"])
        assert {row["gosterge"]: row["tutar_kurus"] for row in summary} == expected["summary_cents"]
        assert {row["gosterge"]: row["tutar_tl"] for row in summary} == expected["summary_tl"]
        customers = sheet_rows(workbook["Müşteri Özeti"])
        for field, metric in (("tutar_kurus", "Toplam fatura"), ("toplam_odeme_kurus", "Eşleşen ödeme"), ("kalan_borc_kurus", "Kalan borç"), ("gecikmis_alacak_kurus", "Gecikmiş borç"), ("fazla_odeme_kurus", "Fazla ödeme")):
            assert sum(row[field] for row in customers) == expected["summary_cents"][metric]
        assert_rows(sheet_rows(workbook["Faturalar"]), expected["invoices"])
        assert all(cell.data_type == "s" for cell in list(workbook["Faturalar"].columns)[0][1:])
        assert_rows(sheet_rows(workbook["Eşleşmeyen Ödemeler"]), expected["unmatched"])
        assert_rows(sheet_rows(workbook["Gelecek Ödemeler"]), expected["future_payments"])
        assert sheet_rows(workbook["Gelecek Faturalar"]) == []
        matched = sheet_rows(workbook["Eşleşen Ödemeler"])
        assert [row["odeme_id"] for row in matched] == expected["matched_payment_ids"]
        assert sum(row["tutar_kurus"] for row in matched) == expected["summary_cents"]["Eşleşen ödeme"]
        assert {row["gecikme_grubu"]: row["borc_kurus"] for row in sheet_rows(workbook["Gecikme Dağılımı"])} == expected["aging_cents"]
    finally:
        workbook.close()

    return {
        "input_format": extension,
        "report_date": report.rapor_tarihi.isoformat(),
        "summary_cents": summarize(report.faturalar),
        "invoice_ids": [row["fatura_no"] for row in report.faturalar],
        "matched_payment_ids": [row["odeme_id"] for row in report.eslesen_odemeler],
        "unmatched_count": len(report.eslesmeyen_odemeler),
        "unmatched_cents": sum(row["tutar_kurus"] for row in report.eslesmeyen_odemeler),
        "future_payment_count": len(report.gelecek_odemeler),
        "future_payment_cents": sum(row["tutar_kurus"] for row in report.gelecek_odemeler),
        "filtered_invoice_ids": filtered["fatura_no"].tolist(),
        "checks": "Sabit beklentiler, dosyaya yazma ve yeniden okuma başarılı.",
        "browser_upload_tested": False,
    }


def main():
    parser = argparse.ArgumentParser(description="Bağımsız CSV/XLSX kabul kontrolü")
    parser.add_argument("--output", type=Path, default=DATA / "reports")
    parser.add_argument("--date", choices=["2026-06-30", "2026-09-26"], default="2026-06-30")
    args = parser.parse_args()
    outcomes = [verify_case(extension, args.output / f"{extension}_input", args.date) for extension in ("csv", "xlsx")]
    assert {key: value for key, value in outcomes[0].items() if key != "input_format"} == {key: value for key, value in outcomes[1].items() if key != "input_format"}
    (args.output / "verification.json").write_text(json.dumps(outcomes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("CSV ve XLSX: sabit beklenen sonuçların tümü doğrulandı.")
    print("Üretilen CSV ve Excel raporları yeniden okunarak doğrulandı.")
    print(f"Rapor klasörü: {args.output.resolve()}")
    print("Tarayıcıdan dosya yükleme denenmedi; bu kontrol dosya işleme/dışa aktarma kabul testidir.")


if __name__ == "__main__":
    main()
