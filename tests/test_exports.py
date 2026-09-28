from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import load_workbook
import pandas as pd
import pytest

from eslestirici.demo import DEMO_DATE, demo_bytes, sample_frame
from eslestirici.exports import csv_bytes, frame, report_sheets, safe_text, xlsx_bytes
from eslestirici.reconciliation import process_files
from eslestirici.money import MAX_AMOUNT_CENTS


@pytest.mark.parametrize("value", ["=SUM(1,2)", "+cmd", "-1+2", "@SUM(A1)", "  =1+1", "\t=1+1", "\r=1", "\n+2", "\ufeff=1+1"])
def test_formula_injection_escaped_in_csv_and_excel(value):
    data = pd.DataFrame([{"musteri": value, "fatura_no": "0001", "tutar_kurus": 125050}], dtype=object)
    parsed = pd.read_csv(StringIO(csv_bytes(data).decode("utf-8-sig")), dtype=str, keep_default_na=False)
    assert parsed.iloc[0]["musteri"].startswith("'")
    assert parsed.iloc[0]["fatura_no"] == "0001"
    assert parsed.iloc[0]["tutar_tl"] == "1250.50"
    workbook = load_workbook(BytesIO(xlsx_bytes({"Faturalar": data})), data_only=False)
    assert workbook.active["A2"].data_type == "s"
    assert workbook.active["A2"].value.startswith("'")
    assert workbook.active["B2"].value == "0001"
    assert workbook.active["B2"].data_type == "s"
    assert workbook.active["C2"].value == 1250.50
    assert workbook.active["C2"].data_type == "n"
    assert workbook.active["C2"].number_format == '#,##0.00 "TL"'
    assert workbook.active["D2"].value == 125050
    assert workbook.active.column_dimensions["D"].hidden
    workbook.close()


def test_safe_normal_text_and_numbers():
    assert safe_text("Hayalî AŞ") == "Hayalî AŞ"
    assert safe_text("0001") == "0001"
    assert safe_text(100) == 100


def test_excel_report_has_all_sheets_and_report_date():
    invoices, payments = demo_bytes()
    report, errors = process_files(invoices, "a.csv", payments, "b.csv", DEMO_DATE)
    assert not errors
    workbook = load_workbook(BytesIO(xlsx_bytes(report_sheets(report))))
    assert len(workbook.sheetnames) == 9
    assert workbook.sheetnames[0] == "Özet"
    assert workbook["Rapor Bilgisi"]["B2"].value == datetime(2026, 6, 30)
    assert workbook["Özet"]["B1"].value == datetime(2026, 6, 30)
    assert workbook["Özet"]["B1"].number_format == "dd.mm.yyyy"
    assert workbook["Özet"].freeze_panes == "A4"
    assert workbook["Özet"].auto_filter.ref == "A3:C8"
    assert workbook["Faturalar"].max_row == 10
    assert workbook["Eşleşmeyen Ödemeler"].max_row == 3
    assert workbook["Faturalar"].freeze_panes == "A2"
    assert workbook["Faturalar"].auto_filter.ref == "A1:O10"
    assert workbook["Faturalar"]["C2"].data_type == "d"
    assert workbook["Faturalar"]["C2"].number_format == "dd.mm.yyyy"
    assert workbook["Faturalar"].row_dimensions[1].height >= 30
    assert [cell.value for cell in workbook["Faturalar"][1]][4:8] == [
        "Fatura tutarı (TL)", "Ödenen (TL)", "Kalan borç (TL)", "Fazla ödeme (TL)",
    ]
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_numeric_money_round_trip_keeps_fractional_cents_and_numeric_sort():
    amounts = [100001, 99999, 1, 29, 57, 0, MAX_AMOUNT_CENTS, MAX_AMOUNT_CENTS * 100]
    data = frame([{"fatura_no": f"{index:04d}", "tutar_kurus": cents} for index, cents in enumerate(amounts)], ["fatura_no", "tutar_kurus"])
    workbook = load_workbook(BytesIO(xlsx_bytes({"Faturalar": data})))
    sheet = workbook.active
    values = []
    for row, cents in zip(sheet.iter_rows(min_row=2), amounts):
        identifier, money, technical = row
        assert identifier.data_type == "s" and identifier.number_format == "@"
        assert money.data_type == technical.data_type == "n"
        assert Decimal(str(money.value)) * 100 == technical.value == cents
        values.append(money.value)
    assert sorted(values[:6]) == [0, 0.01, 0.29, 0.57, 999.99, 1000.01]
    assert sum(Decimal(str(value)) for value in values) * 100 == sum(amounts)
    workbook.close()


def test_excel_refuses_silent_precision_loss_without_changing_csv():
    cents = 9_999_999_999_999_999
    data = frame([{"tutar_kurus": cents}], ["tutar_kurus"])
    with pytest.raises(ValueError, match="hassasiyet.*CSV"):
        xlsx_bytes({"Özet": data})
    assert csv_bytes(data).decode("utf-8-sig").endswith("9999999999999999,99999999999999.99\r\n")


def test_xlsx_demo_round_trip():
    invoices = xlsx_bytes({"Faturalar": sample_frame("faturalar")})
    payments = xlsx_bytes({"Ödemeler": sample_frame("odemeler")})
    excel_report, errors = process_files(invoices, "faturalar.xlsx", payments, "odemeler.xlsx", DEMO_DATE)
    csv_invoices, csv_payments = demo_bytes()
    csv_report, _ = process_files(csv_invoices, "a.csv", csv_payments, "b.csv", DEMO_DATE)
    assert not errors
    assert excel_report == csv_report


def test_empty_exports_keep_headers():
    data = frame([], ["fatura_no", "tutar_kurus"])
    assert csv_bytes(data).decode("utf-8-sig").strip() == "fatura_no,tutar_kurus,tutar_tl"
    workbook = load_workbook(BytesIO(xlsx_bytes({"Boş": data})))
    assert workbook.active.max_row == 1
    assert [cell.value for cell in workbook.active[1]] == ["Fatura no", "Tutar (TL)", "tutar_kurus"]
    assert workbook.active.column_dimensions["C"].hidden
    workbook.close()
