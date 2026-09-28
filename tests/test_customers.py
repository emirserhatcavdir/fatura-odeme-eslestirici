from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
import pytest

from eslestirici.customers import CUSTOMER_COLUMNS, customer_key, customer_summary
from eslestirici.exports import csv_bytes, frame, report_sheets, xlsx_bytes
from eslestirici.models import Fatura, Odeme
from eslestirici.reconciliation import invoice_columns, process_files, reconcile, summarize
from eslestirici.ui import customer_options, customer_summary_html
from eslestirici.views import filter_invoices

AS_OF = date(2026, 6, 30)


@pytest.fixture
def report():
    def invoice(number, name, cents, due, identifier=None, issued=date(2026, 1, 1)):
        return Fatura(number, name, issued, due, cents, identifier)

    return reconcile([
        invoice("1", "A", 10001, date(2026, 6, 10), "001"),
        invoice("2", "A yeni", 5000, date(2026, 6, 20), "001"),
        invoice("3", "A", 2000, date(2026, 1, 2), "001"),  # Gecikerek kapanmış.
        invoice("4", "A", 3000, AS_OF, "001"),  # Vade bugün, açık ama gecikmemiş.
        invoice("5", "A", 99999, date(2026, 7, 1), "001", date(2026, 7, 1)),
        invoice("6", "A", 4000, date(2026, 6, 29), "002"),  # Aynı ad, başka kimlik.
        invoice("7", "A", 2000, date(2026, 7, 1)),  # Kimlikli A'ya atanmaz.
        invoice("8", "a", 1000, date(2026, 7, 1)),  # Esnek arama birleşme değildir.
        invoice("9", "Kapalı", 1000, date(2026, 1, 2)),
        invoice("10", "Gelecek müşteri", 100, date(2026, 7, 1), issued=date(2026, 7, 1)),
    ], [
        Odeme("p1", "1", date(2026, 6, 11), 2000),
        Odeme("p2", "1", AS_OF, 1001),
        Odeme("p3", "2", AS_OF, 7000),  # Fazla ödeme 2000 kuruş.
        Odeme("p4", "3", AS_OF, 2000),
        Odeme("p5", "1", date(2026, 7, 1), 7000),  # Gelecek ödeme.
        Odeme("p6", "5", AS_OF, 99999),  # Gelecek faturaya ödeme eşleşmez.
        Odeme("p7", "olmayan", AS_OF, 30000),
        Odeme("p8", "9", AS_OF, 1000),
    ], AS_OF)


def test_customer_totals_partial_excess_closed_overdue_and_future(report):
    before = deepcopy(report)
    summary = customer_summary(report.faturalar)
    assert summary[0] == {
        "musteri": "A / A yeni", "musteri_id": "001", "tutar_kurus": 20001,
        "toplam_odeme_kurus": 12001, "kalan_borc_kurus": 10000,
        "gecikmis_alacak_kurus": 7000, "fazla_odeme_kurus": 2000,
        "acik_fatura_sayisi": 2, "en_eski_gecikme_gun": 20,
    }
    assert [row["gecikmis_alacak_kurus"] for row in summary] == [7000, 4000, 0, 0, 0]
    closed = next(row for row in summary if row["musteri"] == "Kapalı")
    assert closed["en_eski_gecikme_gun"] == closed["acik_fatura_sayisi"] == 0
    assert all(row["musteri"] != "Gelecek müşteri" for row in summary)
    assert report == before


@pytest.mark.parametrize("filters", [
    {}, {"statuses": ["Kısmen ödendi"]}, {"buckets": ["1–30 gün"]},
    {"query": "a"}, {"customer_keys": [("id", "001")]},
    {"customer_keys": [("name", "A")]},
    {"query": "a", "statuses": ["Ödenmedi"], "customer_keys": [("id", "001")]},
    {"query": "bulunmayan"},
])
def test_summary_and_invoice_scope_reconcile_exactly(report, filters):
    data = frame(report.faturalar, invoice_columns(report.faturalar))
    rows = filter_invoices(data, **filters).to_dict("records")
    groups = customer_summary(rows)
    totals = summarize(rows)
    for field, metric in (("tutar_kurus", "Toplam fatura"), ("toplam_odeme_kurus", "Eşleşen ödeme"), ("kalan_borc_kurus", "Kalan borç"), ("gecikmis_alacak_kurus", "Gecikmiş borç"), ("fazla_odeme_kurus", "Fazla ödeme")):
        assert sum(row[field] for row in groups) == totals[metric]
    for group in groups:
        details = [row for row in rows if customer_key(row) == customer_key(group)]
        assert group["tutar_kurus"] == sum(row["tutar_kurus"] for row in details)
        assert group["acik_fatura_sayisi"] == sum(row["kalan_borc_kurus"] > 0 for row in details)
        assert group["en_eski_gecikme_gun"] == max((row["gecikme_gun"] for row in details if row["kalan_borc_kurus"] > 0 and row["gecikme_gun"] > 0), default=0)


def test_customer_identity_is_exact_and_missing_id_is_not_guessed(report):
    data = frame(report.faturalar, invoice_columns(report.faturalar))
    expected = {("id", "001"): ["1", "2", "3", "4"], ("id", "002"): ["6"], ("name", "A"): ["7"], ("name", "a"): ["8"]}
    for key, invoice_ids in expected.items():
        assert filter_invoices(data, customer_keys=[key])["fatura_no"].tolist() == invoice_ids
    assert len(customer_options(report.faturalar)) == 5


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_optional_customer_id_survives_input_output_and_preserves_leading_zeroes(extension):
    data = pd.DataFrame([
        ["01", "Çınar", "2026-06-01", "2026-06-10", "0.29", "001"],
        ["02", "Çınar", "2026-06-01", "2026-06-10", "0.57", "1"],
        ["03", "Çınar", "2026-06-01", "2026-06-10", "0.01", None],
        ["04", "Cinar", "2026-06-01", "2026-06-10", "0.02", None],
    ], columns=["fatura_no", "musteri", "fatura_tarihi", "vade_tarihi", "tutar", "musteri_id"])
    content = csv_bytes(data) if extension == "csv" else xlsx_bytes({"Faturalar": data})
    result, errors = process_files(content, "faturalar." + extension, b"odeme_id,fatura_no,odeme_tarihi,tutar\n", "p.csv", AS_OF)
    assert not errors
    groups = customer_summary(result.faturalar)
    assert len(groups) == 4
    assert {customer_key(row) for row in groups} == {("id", "001"), ("id", "1"), ("name", "Çınar"), ("name", "Cinar")}
    exported = pd.read_csv(BytesIO(csv_bytes(frame(result.faturalar, invoice_columns(result.faturalar)))), dtype=str, keep_default_na=False)
    assert exported["musteri_id"].tolist() == ["001", "1", "", ""]
    workbook = load_workbook(BytesIO(xlsx_bytes(report_sheets(result))))
    sheet = workbook["Müşteri Özeti"]
    assert sheet["B2"].value == "1" and sheet["B3"].value == "001"
    assert sheet["B3"].data_type == "s" and sheet["B3"].number_format == "@"
    workbook.close()


def test_customer_excel_all_rows_numeric_money_and_empty_schema(report):
    sheets = report_sheets(report)
    workbook = load_workbook(BytesIO(xlsx_bytes(sheets)))
    sheet = workbook["Müşteri Özeti"]
    assert sheet.max_row == 6 and sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:N6"
    for cells, expected in zip(sheet.iter_rows(min_row=2), customer_summary(report.faturalar)):
        for offset, key in enumerate(CUSTOMER_COLUMNS[2:7]):
            money, cents = cells[2 + offset], cells[9 + offset]
            assert money.data_type == cents.data_type == "n"
            assert Decimal(str(money.value)) * 100 == cents.value == expected[key]
            assert money.number_format == '#,##0.00 "TL"'
            assert sheet.column_dimensions[cents.column_letter].hidden
        assert cells[7].data_type == cells[8].data_type == "n"
    workbook.close()
    empty = reconcile([], [], AS_OF)
    assert report_sheets(empty)["Müşteri Özeti"].columns.tolist() == CUSTOMER_COLUMNS
    assert customer_summary([]) == []


def test_large_customer_totals_use_python_integer_cents():
    # Tek müşteri toplamı int64 sınırını aşsa da toplama Python int ile kesindir.
    row = {"musteri": "A", "tutar_kurus": 10**18 + 1, "toplam_odeme_kurus": 1,
           "kalan_borc_kurus": 10**18, "fazla_odeme_kurus": 0, "gecikme_gun": 1}
    summary = customer_summary([row] * 10)[0]
    assert summary["tutar_kurus"] == 10**19 + 10
    assert summary["kalan_borc_kurus"] == 10**19


def test_customer_html_escapes_names_and_ids(report):
    rows = customer_summary(report.faturalar)
    rows[0]["musteri"] = '<img src=x onerror="bad()">'
    rows[0]["musteri_id"] = "<script>bad()</script>"
    html = customer_summary_html(rows)
    assert "<script>" not in html and "<img" not in html
    assert "&lt;img" in html and "&lt;script&gt;" in html
    assert 'data-label="En eski gecikme (gün)"' in html


def test_closed_invoice_with_historical_delay_does_not_inflate_oldest(report):
    rows = deepcopy(report.faturalar)
    next(row for row in rows if row["fatura_no"] == "3")["gecikme_gun"] = 900
    assert customer_summary(rows)[0]["en_eski_gecikme_gun"] == 20


def test_optional_customer_id_validation():
    content = 'fatura_no,musteri,fatura_tarihi,vade_tarihi,tutar,musteri_id\n1,A,2026-06-01,2026-06-02,1.00,"bad\nid"\n'.encode()
    report, errors = process_files(content, "i.csv", b"odeme_id,fatura_no,odeme_tarihi,tutar\n", "p.csv", AS_OF)
    assert report is None
    assert len(errors) == 1 and errors[0].alan == "musteri_id"
