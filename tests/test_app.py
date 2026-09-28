from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from streamlit.dataframe_util import convert_arrow_bytes_to_pandas_df
from streamlit.testing.v1 import AppTest

from eslestirici.demo import DEMO_DATE, demo_bytes, sample_frame
from eslestirici.exports import EXCEL_EXPORT_VERSION, csv_bytes, xlsx_bytes

APP = str(Path(__file__).resolve().parents[1] / "sap.py")


def start_demo():
    app = AppTest.from_file(APP, default_timeout=30).run()
    assert not app.exception
    app.button(key="sidebar_demo").click().run()
    assert not app.exception
    return app


def metrics(app):
    return {metric.label: metric.value for metric in app.metric}


def test_initial_screen_and_demo():
    app = start_demo()
    assert app.title[0].value == "Fatura–Ödeme Eşleştirici"
    assert app.date_input[0].value == DEMO_DATE
    assert metrics(app) == {"Toplam fatura": "19.050,75 TL", "Eşleşen ödeme": "7.950,50 TL", "Kalan borç": "11.300,25 TL", "Gecikmiş borç": "9.900,00 TL", "Fazla ödeme": "200,00 TL"}
    assert len(app.dataframe[0].value) == 9


def test_main_demo_action_sets_date_and_places_current_summary_above_search():
    app = AppTest.from_file(APP, default_timeout=30).run()
    assert not app.exception and not app.metric
    assert app.main.button(key="main_demo").label == "Örnek raporu incele"
    app.main.button(key="main_demo").click().run()
    assert not app.exception
    assert app.date_input[0].value == DEMO_DATE
    assert app.radio[0].value == "Örnek veriler"
    assert metrics(app)["Toplam fatura"] == "19.050,75 TL"
    elements = list(app.main)
    assert next(i for i, item in enumerate(elements) if item.type == "metric") < next(i for i, item in enumerate(elements) if item.type == "text_input")
    assert not any("Filtre uygulanıyor" in item.value for item in app.caption)


def test_filters_change_cards_and_invoices_but_not_excel():
    app = start_demo()
    excel_before = app.session_state["report_bundle"]["excel"]
    app.multiselect(key="status_filter").set_value(["Kısmen ödendi"]).run()
    assert not app.exception
    assert metrics(app)["Kalan borç"] == "2.000,00 TL"
    assert len(app.dataframe[0].value) == 1
    assert any("Filtre uygulanıyor · 1 / 9 fatura" in item.value for item in app.caption)
    assert app.session_state["report_bundle"]["excel"] == excel_before
    app.multiselect(key="customer_filter").set_value([("name", "Hayalî Ada Kitap")]).run()
    assert not app.exception
    assert metrics(app)["Toplam fatura"] == "0,00 TL"


def test_aging_filter_and_date_change():
    app = start_demo()
    app.multiselect(key="aging_filter").set_value(["90 üzeri gün"]).run()
    assert not app.exception
    assert metrics(app)["Gecikmiş borç"] == "4.200,00 TL"
    app.date_input[0].set_value(date(2026, 7, 1)).run()
    assert not app.exception
    assert len(app.dataframe[0].value) == 10
    assert app.multiselect(key="aging_filter").value == []
    assert metrics(app)["Eşleşen ödeme"] == "9.050,75 TL"


def test_switch_from_demo_to_upload_clears_results():
    app = start_demo()
    app.radio[0].set_value("Dosya yükle").run()
    assert not app.exception
    assert len(app.metric) == 0
    assert "report_bundle" not in app.session_state


class Upload(BytesIO):
    def __init__(self, content, name):
        super().__init__(content)
        self.name = name


def upload_app(invoice_content, payment_content, extension="csv"):
    def uploader(label, **kwargs):
        if kwargs["key"] == "invoice_upload":
            return Upload(invoice_content, "faturalar." + extension)
        return Upload(payment_content, "odemeler." + extension)

    app = AppTest.from_file(APP, default_timeout=30)
    app.session_state["report_date"] = DEMO_DATE
    # AppTest'te dosya seçme diyaloğu yerine yükleme nesnesi taklit edilir.
    with patch("streamlit.file_uploader", side_effect=uploader):
        app.run()
    return app


def test_uploaded_csv_report():
    app = upload_app(*demo_bytes())
    assert not app.exception
    assert metrics(app)["Toplam fatura"] == "19.050,75 TL"


def test_uploaded_xlsx_report():
    invoices = xlsx_bytes({"Faturalar": sample_frame("faturalar")})
    payments = xlsx_bytes({"Ödemeler": sample_frame("odemeler")})
    app = upload_app(invoices, payments, "xlsx")
    assert not app.exception
    assert metrics(app)["Kalan borç"] == "11.300,25 TL"


def test_invalid_upload_shows_errors_without_final_report():
    invoices, payments = demo_bytes()
    app = upload_app(invoices, payments + b"P0001,0001,2026-06-01,1.00\n")
    assert not app.exception
    assert len(app.error) == 1
    assert "Nihai rapor oluşturulmadı" in app.error[0].value
    assert len(app.metric) == 0
    assert "Satır" in app.dataframe[0].value.columns


def test_header_only_uploads_are_empty_report():
    app = upload_app(csv_bytes(sample_frame("faturalar", True)), csv_bytes(sample_frame("odemeler", True)))
    assert not app.exception
    assert all(metric.value == "0,00 TL" for metric in app.metric)


def test_zero_byte_upload_is_actionable_error():
    _, payments = demo_bytes()
    app = upload_app(b"", payments)
    assert not app.exception
    assert len(app.error) == 1
    assert len(app.metric) == 0


def test_search_updates_cards_table_csv_but_keeps_full_excel():
    app = start_demo()
    excel_before = app.session_state["report_bundle"]["excel"]
    with patch("streamlit.download_button", wraps=st.download_button) as downloads:
        app.text_input(key="invoice_search").set_value("0001").run()
    assert not app.exception
    assert metrics(app)["Kalan borç"] == "2.000,00 TL"
    assert app.dataframe[0].value["Fatura no"].tolist() == ["0001"]
    assert any("Filtre uygulanıyor · 1 / 9 fatura" in item.value for item in app.caption)
    csv_call = next(call for call in downloads.call_args_list if call.args[0] == "Filtreli faturaları CSV indir")
    downloaded = pd.read_csv(BytesIO(csv_call.args[1]), dtype=str)
    assert downloaded["fatura_no"].tolist() == ["0001"]
    assert downloaded.iloc[0]["kalan_borc_tl"] == "2000.00"
    assert "fazla_odeme_kurus" in downloaded.columns
    assert app.session_state["report_bundle"]["excel"] == excel_before
    assert "2 eşleşmeyen ödeme" in app.warning[0].value
    assert "850,75 TL" in app.warning[0].value
    assert "#eslesmeyen-odemeler" in app.warning[0].value


def test_clear_filters_preserves_uploaded_files_reporting_date_and_excel():
    invoices, payments = demo_bytes()

    def uploader(label, **kwargs):
        return Upload(invoices, "faturalar.csv") if kwargs["key"] == "invoice_upload" else Upload(payments, "odemeler.csv")

    app = AppTest.from_file(APP, default_timeout=30)
    app.session_state["report_date"] = DEMO_DATE
    with patch("streamlit.file_uploader", side_effect=uploader):
        app.run()
        initial_bundle = app.session_state["report_bundle"]
        app.text_input(key="invoice_search").set_value("mavi")
        app.multiselect(key="customer_filter").set_value([("name", "Hayalî Mavi Kırtasiye")])
        app.multiselect(key="status_filter").set_value(["Kısmen ödendi"])
        app.multiselect(key="aging_filter").set_value(["1–30 gün"]).run()
        app.selectbox(key="selected_invoice").set_value("0001").run()
        app.button(key="clear_filters").click().run()
    assert not app.exception
    assert app.radio[0].value == "Dosya yükle"
    assert app.date_input[0].value == DEMO_DATE
    assert app.text_input(key="invoice_search").value == ""
    assert all(widget.value == [] for widget in app.multiselect)
    assert app.selectbox(key="selected_invoice").value is None
    assert app.session_state["report_bundle"]["signature"] == initial_bundle["signature"]
    assert app.session_state["report_bundle"]["excel"] == initial_bundle["excel"]
    assert len(app.dataframe[0].value) == 9
    assert not any("Filtre uygulanıyor" in item.value for item in app.caption)


def test_old_cached_excel_is_replaced_without_resetting_report_or_search():
    app = start_demo()
    app.text_input(key="invoice_search").set_value("mavi").run()
    bundle = app.session_state["report_bundle"]
    signature = bundle["signature"]
    bundle["excel"] = b"old-text-money-export"
    bundle.pop("excel_version", None)
    app.run()
    assert not app.exception
    assert app.text_input(key="invoice_search").value == "mavi"
    assert app.session_state["report_bundle"]["signature"] == signature
    assert app.session_state["report_bundle"]["excel"].startswith(b"PK")
    assert app.session_state["report_bundle"]["excel_version"] == EXCEL_EXPORT_VERSION


def test_excel_precision_error_preserves_csv_download_and_report():
    from eslestirici.exports import ExcelPrecisionError

    app = start_demo()
    app.session_state["report_bundle"].pop("excel_version", None)
    with patch("eslestirici.exports.xlsx_bytes", side_effect=ExcelPrecisionError("Kuruş hassasiyeti için CSV raporunu kullanın.")):
        app.run()
    assert not app.exception
    assert any("CSV raporunu" in item.value for item in app.error)
    assert any(item.key == "filtered_csv" for item in app.get("download_button"))
    assert not any(item.key == "full_excel" for item in app.get("download_button"))
    assert metrics(app)["Toplam fatura"] == "19.050,75 TL"


def test_detail_shows_installments_for_selected_invoice_only():
    app = start_demo()
    app.selectbox(key="selected_invoice").set_value("0005").run()
    assert not app.exception
    table = next(item for item in app.dataframe if "Ödeme kimliği" in item.value and item.value["Ödeme kimliği"].tolist() == ["P0004", "P0005"])
    assert table.value["Tutar (TL)"].tolist() == ["1.000,00 TL", "1.500,00 TL"]
    assert any("Fatura no: 0005" in item.value for item in app.text)
    assert any("Toplam ödeme" == item.value for item in app.caption)
    assert metrics(app)["Eşleşen ödeme"] == "7.950,50 TL"


def test_detail_future_payment_is_separate_and_does_not_change_totals():
    app = start_demo()
    before = metrics(app)
    app.selectbox(key="selected_invoice").set_value("0007").run()
    assert not app.exception
    assert metrics(app) == before
    assert any("raporlama tarihine kadar gerçekleşmiş ödemesi yok" in item.value for item in app.info)
    assert any("gelecek tarihli ödemeler toplam ödemeye" in item.value for item in app.info)
    future_tables = [item.value for item in app.dataframe if "Ödeme kimliği" in item.value and item.value["Ödeme kimliği"].tolist() == ["P0007"]]
    assert future_tables
    assert all(table["Tutar (TL)"].tolist() == ["600,25 TL"] for table in future_tables)
    assert any(item.value == "0,00 TL" for item in app.markdown)


def test_unpaid_invoice_has_clear_empty_payment_messages():
    app = start_demo()
    app.selectbox(key="selected_invoice").set_value("0003").run()
    assert not app.exception
    assert any("raporlama tarihine kadar gerçekleşmiş ödemesi yok" in item.value for item in app.info)
    assert any("gelecek tarihli ödeme yok" in item.value for item in app.caption)


def test_filter_changes_clear_detail_even_if_invoice_is_still_in_results():
    app = start_demo()
    app.selectbox(key="selected_invoice").set_value("0001").run()
    app.multiselect(key="status_filter").set_value(["Kısmen ödendi"]).run()
    assert not app.exception
    assert app.dataframe[0].value["Fatura no"].tolist() == ["0001"]
    assert app.selectbox(key="selected_invoice").value is None
    assert not any(item.value == "Seçili fatura" for item in app.subheader)
    app.selectbox(key="selected_invoice").set_value("0001").run()
    app.text_input(key="invoice_search").set_value("0001").run()
    assert app.selectbox(key="selected_invoice").value is None


def test_source_change_clears_detail_even_when_uploaded_bytes_match_demo():
    app = start_demo()
    app.selectbox(key="selected_invoice").set_value("0001").run()
    invoices, payments = demo_bytes()

    def uploader(label, **kwargs):
        return Upload(invoices, "ornek_faturalar.csv") if kwargs["key"] == "invoice_upload" else Upload(payments, "ornek_odemeler.csv")

    with patch("streamlit.file_uploader", side_effect=uploader):
        app.radio[0].set_value("Dosya yükle").run()
    assert not app.exception
    assert app.selectbox(key="selected_invoice").value is None
    assert not any(item.value == "Seçili fatura" for item in app.subheader)


def test_new_file_contents_clear_detail_in_same_upload_source():
    invoices, payments = demo_bytes()
    app = upload_app(invoices, payments)

    def uploader(label, **kwargs):
        return Upload(invoices, "faturalar.csv") if kwargs["key"] == "invoice_upload" else Upload(payments, "odemeler.csv")

    with patch("streamlit.file_uploader", side_effect=uploader):
        app.selectbox(key="selected_invoice").set_value("0001").run()
        invoices = invoices.replace(b"5000.00", b"5000.01")
        app.run()
    assert not app.exception
    assert app.selectbox(key="selected_invoice").value is None
    assert metrics(app)["Kalan borç"] == "11.300,26 TL"


def test_date_change_clears_detail_and_includes_payment_on_that_date():
    app = start_demo()
    app.selectbox(key="selected_invoice").set_value("0007").run()
    app.date_input[0].set_value(date(2026, 7, 1)).run()
    assert not app.exception
    assert app.selectbox(key="selected_invoice").value is None
    app.selectbox(key="selected_invoice").set_value("0007").run()
    assert not any("gelecek tarihli ödemeler toplam ödemeye" in item.value for item in app.info)
    assert any("gelecek tarihli ödeme yok" in item.value for item in app.caption)
    assert metrics(app)["Eşleşen ödeme"] == "9.050,75 TL"


def test_numeric_source_turkish_format_status_colors_and_column_toggle():
    app = start_demo()
    grid = app.dataframe[0]
    assert grid.value.columns.tolist() == ["Fatura no", "Müşteri", "Vade", "Fatura tutarı", "Ödenen", "Kalan", "Durum"]
    assert pd.api.types.is_integer_dtype(grid.value["Kalan"])
    display = convert_arrow_bytes_to_pandas_df(grid.proto.arrow_data.styler.display_values)
    assert display.iloc[0]["Kalan"] == "2.000,00 TL"
    assert display.iloc[0]["Durum"] == "Kısmen ödendi"
    assert "background-color" in grid.proto.arrow_data.styler.styles
    app.selectbox(key="selected_invoice").set_value("0001").run()
    app.checkbox(key="detailed_columns").check().run()
    assert not app.exception
    assert len(app.dataframe[0].value.columns) == 11
    assert app.selectbox(key="selected_invoice").value == "0001"


def test_no_search_results_hide_old_detail_and_show_empty_state():
    app = start_demo()
    app.selectbox(key="selected_invoice").set_value("0001").run()
    app.text_input(key="invoice_search").set_value("bulunmayan müşteri").run()
    assert not app.exception
    assert len(app.selectbox) == 0
    assert all(value == "0,00 TL" for value in metrics(app).values())
    assert any("gösterilecek fatura yok" in item.value for item in app.info)
    assert any("gecikmiş borç yok" in item.value for item in app.success)
    assert not any(item.value == "Seçili fatura" for item in app.subheader)


def test_empty_reporting_date_does_not_show_previous_report():
    app = start_demo()
    app.date_input[0].set_value(None).run()
    assert not app.exception
    assert not app.metric
    assert any("bir raporlama tarihi seç" in item.value for item in app.info)


def customer_html(app):
    return next(item.proto.body for item in app.get("html") if 'class="customer-summary"' in item.proto.body)


def test_customer_selection_clear_and_full_excel_scope():
    app = start_demo()
    initial_excel = app.session_state["report_bundle"]["excel"]
    html = customer_html(app)
    assert html.index("Hayalî Lale Tasarım") < html.index("Hayalî Ufuk Seramik") < html.index("Hayalî Mavi Kırtasiye")
    app.multiselect(key="customer_filter").set_value([("name", "Hayalî Ada Kitap")]).run()
    assert not app.exception
    assert app.dataframe[0].value["Fatura no"].tolist() == ["0004", "0007"]
    html = customer_html(app)
    assert "Hayalî Ada Kitap" in html and "Hayalî Mavi Kırtasiye" not in html
    assert 'data-label="Kalan alacak">600,25 TL' in html
    assert 'data-label="Fazla ödeme">200,00 TL' in html
    assert 'data-label="Açık fatura">1' in html
    assert 'data-label="En eski gecikme (gün)">0' in html
    assert app.session_state["report_bundle"]["excel"] == initial_excel
    workbook = load_workbook(BytesIO(initial_excel))
    assert workbook["Müşteri Özeti"].max_row == 6
    assert workbook["Faturalar"].max_row == 10
    workbook.close()
    app.selectbox(key="selected_invoice").set_value("0004").run()
    app.button(key="clear_customer").click().run()
    assert not app.exception
    assert app.multiselect(key="customer_filter").value == []
    assert app.selectbox(key="selected_invoice").value is None
    assert len(app.dataframe[0].value) == 9
    assert app.date_input[0].value == DEMO_DATE
    assert customer_html(app).count('<th scope="row">') == 5


def test_clear_customer_keeps_search_status_and_aging_filters():
    app = start_demo()
    app.text_input(key="invoice_search").set_value("hayali")
    app.multiselect(key="status_filter").set_value(["Ödenmedi"])
    app.multiselect(key="aging_filter").set_value(["1–30 gün", "31–60 gün", "61–90 gün", "90 üzeri gün"])
    app.multiselect(key="customer_filter").set_value([("name", "Hayalî Ufuk Seramik")]).run()
    assert app.dataframe[0].value["Fatura no"].tolist() == ["0009", "0010"]
    assert 'data-label="Gecikmiş alacak">3.700,00 TL' in customer_html(app)
    app.button(key="clear_customer").click().run()
    assert not app.exception
    assert app.text_input(key="invoice_search").value == "hayali"
    assert app.multiselect(key="status_filter").value == ["Ödenmedi"]
    assert len(app.multiselect(key="aging_filter").value) == 4
    assert app.dataframe[0].value["Fatura no"].tolist() == ["0003", "0009", "0010"]
    assert customer_html(app).count('<th scope="row">') == 2


def test_date_change_clears_customer_and_recalculates_summary():
    app = start_demo()
    app.multiselect(key="customer_filter").set_value([("name", "Hayalî Ada Kitap")]).run()
    app.date_input[0].set_value(date(2026, 7, 1)).run()
    assert not app.exception
    assert app.multiselect(key="customer_filter").value == []
    app.multiselect(key="customer_filter").set_value([("name", "Hayalî Ada Kitap")]).run()
    assert 'data-label="Kalan alacak">0,00 TL' in customer_html(app)
    assert 'data-label="Açık fatura">0' in customer_html(app)
    assert 'data-label="Eşleşen ödeme">1.800,25 TL' in customer_html(app)


def test_same_name_customer_ids_select_separate_invoices():
    invoices, payments = demo_bytes()
    data = pd.read_csv(BytesIO(invoices), dtype=str)
    data["musteri_id"] = ["001", None, None, None, None, "002", None, None, None, None]
    invoices = csv_bytes(data)

    def uploader(label, **kwargs):
        return Upload(invoices, "faturalar.csv") if kwargs["key"] == "invoice_upload" else Upload(payments, "odemeler.csv")

    with patch("streamlit.file_uploader", side_effect=uploader):
        app = AppTest.from_file(APP, default_timeout=30)
        app.session_state["report_date"] = DEMO_DATE
        app.run()
        app.multiselect(key="customer_filter").set_value([("id", "002")]).run()
    assert not app.exception
    assert app.dataframe[0].value["Fatura no"].tolist() == ["0006"]
    assert app.dataframe[0].value["Müşteri kimliği"].tolist() == ["002"]
    assert 'data-label="Toplam fatura">800,00 TL' in customer_html(app)
