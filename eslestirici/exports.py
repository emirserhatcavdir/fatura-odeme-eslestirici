from dataclasses import asdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Issue, Report
from .customers import CUSTOMER_COLUMNS, customer_summary
from .money import decimal_amount, format_tl
from .reconciliation import PAYMENT_COLUMNS, aging_totals, invoice_columns, summarize


def safe_text(value: object) -> object:
    """CSV ve Excel'de kullanıcı metinlerinin formül olarak açılmasını engeller."""
    if isinstance(value, str):
        probe = value.lstrip(" \t\r\n\v\f\ufeff")
        if probe.startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def frame(rows: list[dict], columns: list[str]) -> pd.DataFrame:
    # dtype=object kimlikleri ve Python int değerlerini dönüştürmeden saklar.
    result = pd.DataFrame(rows, columns=columns, dtype=object)
    if "musteri_id" in result.columns:
        result["musteri_id"] = result["musteri_id"].where(result["musteri_id"].notna(), None)
    return result


def export_frame(data: pd.DataFrame) -> pd.DataFrame:
    result = data.copy()
    for column in list(result.columns):
        if column.endswith("_kurus"):
            result[column.removesuffix("_kurus") + "_tl"] = result[column].map(decimal_amount)
    return result.map(safe_text)


def csv_bytes(data: pd.DataFrame) -> bytes:
    return export_frame(data).to_csv(index=False, lineterminator="\r\n").encode("utf-8-sig")


EXCEL_MONEY_FORMAT = '#,##0.00 "TL"'
EXCEL_DATE_FORMAT = "dd.mm.yyyy"
EXCEL_EXPORT_VERSION = 3
EXCEL_LABELS = {
    "fatura_no": "Fatura no", "musteri": "Müşteri", "fatura_tarihi": "Fatura tarihi",
    "vade_tarihi": "Vade tarihi", "toplam_odeme_tl": "Ödenen (TL)",
    "kalan_borc_tl": "Kalan borç (TL)", "fazla_odeme_tl": "Fazla ödeme (TL)",
    "borc_tl": "Gecikmiş borç (TL)", "tutar_tl": "Tutar (TL)", "durum": "Durum",
    "gecikme_gun": "Gecikme günü", "gecikme_grubu": "Gecikme grubu",
    "odeme_id": "Ödeme kimliği", "odeme_tarihi": "Ödeme tarihi", "neden": "Açıklama",
    "gosterge": "Gösterge", "alan": "Alan", "deger": "Değer",
    "musteri_id": "Müşteri kimliği", "gecikmis_alacak_tl": "Gecikmiş alacak (TL)",
    "acik_fatura_sayisi": "Açık fatura sayısı", "en_eski_gecikme_gun": "En eski gecikme (gün)",
}


class ExcelPrecisionError(ValueError):
    """Excel'in sayısal sınırında kuruş kaybı oluşacaksa indirmeyi engeller."""


def excel_amount(cents: int) -> Decimal:
    """Yalnızca çıktı sınırında TL'ye çevir; Excel'de sessiz kuruş kaybını engelle."""
    amount = Decimal(int(cents)) / 100
    # Excel 15 anlamlı basamak saklar; büyük toplamlarda bu sınır aşılabilir.
    # openpyxl'in sayısal XML yazımı da aynı kuruşa geri dönmelidir.
    if any(Decimal(format(float(amount), spec)) * 100 != cents for spec in (".15g", ".16g")):
        raise ExcelPrecisionError("Bu toplam Excel'in sayısal hassasiyetini aşıyor; kuruş kaybını önlemek için CSV raporunu kullanın.")
    return amount


def xlsx_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, data in sheets.items():
        sheet = workbook.create_sheet(title)
        money_columns = [column for column in data.columns if column.endswith("_kurus")]
        # Giriş şablonları tutar sütunu kullanır: onların makine şeması değişmez.
        is_report = bool(money_columns) or title == "Rapor Bilgisi"
        columns = [column.removesuffix("_kurus") + "_tl" if column in money_columns else column for column in data.columns]
        columns += money_columns
        labels = dict(EXCEL_LABELS) if is_report else {}
        if title == "Müşteri Özeti":
            labels.update(toplam_odeme_tl="Eşleşen ödeme (TL)", kalan_borc_tl="Kalan alacak (TL)")
        if "musteri" in data.columns and money_columns:
            labels["tutar_tl"] = "Fatura tutarı (TL)"
        elif "odeme_id" in data.columns and money_columns:
            labels["tutar_tl"] = "Ödeme tutarı (TL)"
        header_row = 1
        if title == "Özet" and "report_date" in data.attrs:
            sheet.append(["Raporlama tarihi", data.attrs["report_date"]])
            sheet["A1"].font = Font(name="Arial", size=12, bold=True, color="172B45")
            sheet["B1"].font = Font(name="Arial", size=12, bold=True, color="087F8C")
            sheet["B1"].number_format = EXCEL_DATE_FORMAT
            sheet.row_dimensions[1].height = 28
            header_row = 3
        for index, column in enumerate(columns, 1):
            sheet.cell(header_row, index, labels.get(column, column))
        for row_number, row in enumerate(data.to_dict("records"), header_row + 1):
            for index, column in enumerate(columns, 1):
                cents_column = column.removesuffix("_tl") + "_kurus"
                value = excel_amount(row[cents_column]) if column.endswith("_tl") and cents_column in money_columns else row[column]
                # Kullanıcı metni hiçbir zaman Excel formülü olarak yazılmaz.
                cell = sheet.cell(row_number, index, safe_text(value))
                cell.font = Font(name="Arial", size=11, color="172B45")
                cell.alignment = Alignment(vertical="center")
                if column.endswith("_tl") and cents_column in money_columns:
                    cell.number_format = EXCEL_MONEY_FORMAT
                elif isinstance(value, (date, datetime)):
                    cell.number_format = EXCEL_DATE_FORMAT
                elif isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{sheet.max_row}"
        sheet.print_title_rows = f"1:{header_row}"
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="172B45")
            cell.font = Font(name="Arial", size=11, color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[header_row].height = 34
        for index, column in enumerate(columns, 1):
            letter = get_column_letter(index)
            width = max(len(str(sheet.cell(header_row, index).value)), 14)
            for cells in sheet.iter_rows(min_row=header_row + 1, min_col=index, max_col=index):
                cell = cells[0]
                if isinstance(cell.value, (date, datetime)):
                    length = 10
                elif cell.number_format == EXCEL_MONEY_FORMAT:
                    length = len(f"{cell.value:,.2f} TL")
                else:
                    length = len(str(cell.value if cell.value is not None else ""))
                width = max(width, length)
            sheet.column_dimensions[letter].width = min(width + 4, 52)
            sheet.column_dimensions[letter].hidden = column in money_columns
        if title == "Özet":
            sheet.column_dimensions["A"].width = max(sheet.column_dimensions["A"].width, 25)
            sheet.column_dimensions["B"].width = max(sheet.column_dimensions["B"].width, 24)
        if title == "Rapor Bilgisi":
            sheet.column_dimensions["B"].width = 92
            for row in sheet.iter_rows(min_row=2):
                row[1].alignment = Alignment(wrap_text=True, vertical="center")
                sheet.row_dimensions[row[0].row].height = 34
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def report_sheets(report: Report) -> dict[str, pd.DataFrame]:
    summary = [{"gosterge": name, "tutar_kurus": cents} for name, cents in summarize(report.faturalar).items()]
    info = [
        {"alan": "Raporlama tarihi", "deger": report.rapor_tarihi},
        {"alan": "Kapsam", "deger": "Excel tüm raporu içerir; arayüz filtreleri Excel'i etkilemez."},
        {"alan": "Müşteri özeti", "deger": "Raporlama tarihi kapsamındaki tüm faturalar; gecikmiş alacak azalan sıralıdır. Eşleşmeyen ve gelecek ödemeler dahil değildir."},
        {"alan": "Müşteri gruplama", "deger": "Varsa musteri_id, yoksa birebir müşteri adı kullanılır. Kimliksiz kayıtlar kimlikli müşterilere atanmaz. En eski gecikme yalnızca açık ve gecikmiş faturalardan hesaplanır."},
        {"alan": "Para birimi", "deger": "TL alanları iki ondalık basamaklı sayıdır. Tam sayı *_kurus sütunları korunur ve varsayılan olarak gizlidir."},
        {"alan": "Gecikme", "deger": "Yalnızca kalan borç, vade tarihinden sonraki gün gecikir."},
        {"alan": "Fazla ödeme", "deger": "Başka faturaların borcundan düşülmez."},
        {"alan": "Metin güvenliği", "deger": "Formül başlangıcı içeren metinlere tek tırnak eklenir."},
    ]
    summary_frame = frame(summary, ["gosterge", "tutar_kurus"])
    summary_frame.attrs["report_date"] = report.rapor_tarihi
    return {
        "Özet": summary_frame,
        "Rapor Bilgisi": pd.DataFrame(info),
        "Müşteri Özeti": frame(customer_summary(report.faturalar), CUSTOMER_COLUMNS),
        "Faturalar": frame(report.faturalar, invoice_columns(report.faturalar)),
        "Eşleşen Ödemeler": frame(report.eslesen_odemeler, PAYMENT_COLUMNS),
        "Eşleşmeyen Ödemeler": frame(report.eslesmeyen_odemeler, PAYMENT_COLUMNS + ["neden"]),
        "Gelecek Faturalar": frame(report.gelecek_faturalar, invoice_columns(report.gelecek_faturalar, results=False)),
        "Gelecek Ödemeler": frame(report.gelecek_odemeler, PAYMENT_COLUMNS),
        "Gecikme Dağılımı": frame(aging_totals(report.faturalar), ["gecikme_grubu", "borc_kurus"]),
    }


def issues_frame(issues: list[Issue]) -> pd.DataFrame:
    result = pd.DataFrame([asdict(issue) for issue in issues], columns=["dosya", "satir", "alan", "aciklama"])
    result["satir"] = pd.array(result["satir"], dtype="Int64")
    return result.rename(columns={"dosya": "Dosya", "satir": "Satır", "alan": "Alan", "aciklama": "Açıklama"})


DISPLAY_LABELS = {
    "musteri_id": "Müşteri kimliği",
    "fatura_no": "Fatura no", "musteri": "Müşteri", "fatura_tarihi": "Fatura tarihi", "vade_tarihi": "Vade tarihi",
    "tutar_kurus": "Tutar (TL)", "toplam_odeme_kurus": "Toplam ödeme (TL)", "kalan_borc_kurus": "Kalan borç (TL)",
    "fazla_odeme_kurus": "Fazla ödeme (TL)", "durum": "Durum", "gecikme_gun": "Gecikme (gün)",
    "gecikme_grubu": "Gecikme grubu", "odeme_id": "Ödeme kimliği", "odeme_tarihi": "Ödeme tarihi", "neden": "Açıklama",
}


def display_frame(data: pd.DataFrame) -> pd.DataFrame:
    result = data.copy()
    for column in result.columns:
        if column.endswith("_kurus"):
            result[column] = result[column].map(format_tl)
        elif column.endswith("_tarihi"):
            result[column] = result[column].map(lambda value: value.isoformat())
    return result.rename(columns=DISPLAY_LABELS)
