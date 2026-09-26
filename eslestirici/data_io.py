"""CSV/XLSX okuyucuları; dosyalar bellekte işlenir, diske kaydedilmez."""

import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .models import Issue, Table

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000


def _headers(table: Table, values: list) -> bool:
    table.columns = [str(value).strip() if value is not None else "" for value in values]
    if not table.columns or not any(table.columns):
        table.issues.append(Issue(table.filename, 1, "", "Dosya boş veya sütun başlıkları bulunamadı."))
        return False
    if "" in table.columns:
        table.issues.append(Issue(table.filename, 1, "", "Boş sütun adı var; her sütuna bir başlık verin."))
    if len(set(table.columns)) != len(table.columns):
        table.issues.append(Issue(table.filename, 1, "", "Yinelenen sütun adı var; sütun başlıkları benzersiz olmalı."))
    return not table.issues


def _append_row(table: Table, number: int, values: list) -> bool:
    if all(value is None or str(value).strip() == "" for value in values):
        return True
    if len(table.rows) >= MAX_ROWS:
        table.issues.append(Issue(table.filename, number, "", "Bir dosyada en fazla 50.000 veri satırı desteklenir."))
        return False
    if len(values) != len(table.columns):
        table.issues.append(Issue(table.filename, number, "", "Satırdaki alan sayısı başlıklarla uyuşmuyor. CSV ayırıcı virgül olmalı."))
    else:
        table.rows.append((number, dict(zip(table.columns, values))))
    return True


def read_table(content: bytes, filename: str) -> Table:
    table = Table(filename)
    if len(content) > MAX_BYTES:
        table.issues.append(Issue(filename, None, "", "Dosya boyutu 10 MB sınırını aşıyor."))
        return table
    if not content.strip():
        table.issues.append(Issue(filename, 1, "", "Dosya boş. Sütun başlıklarını içeren şablonu kullanın."))
        return table
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        _read_csv(content, table)
    elif extension == ".xlsx":
        _read_xlsx(content, table)
    else:
        table.issues.append(Issue(filename, None, "", "Yalnızca CSV ve XLSX dosyaları desteklenir."))
    return table


def _read_csv(content: bytes, table: Table) -> None:
    reader = None
    try:
        # utf-8-sig hem BOM içeren hem de BOM içermeyen UTF-8'i okur.
        reader = csv.reader(StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        if not _headers(table, next(reader, [])):
            return
        while True:
            line = reader.line_num + 1
            values = next(reader, None)
            if values is None or not _append_row(table, line, values):
                break
    except UnicodeDecodeError:
        table.issues.append(Issue(table.filename, None, "", "CSV kodlaması UTF-8 olmalı. Dosyayı CSV UTF-8 olarak kaydedin."))
    except csv.Error:
        table.issues.append(Issue(table.filename, reader.line_num if reader else None, "", "CSV yapısı bozuk; tırnakları ve virgül ayırıcılarını kontrol edin."))


def _read_xlsx(content: bytes, table: Table) -> None:
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                table.issues.append(Issue(table.filename, None, "", "XLSX açılmış içerik boyutu 100 MB sınırını aşıyor."))
                return
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if not workbook.worksheets:
            table.issues.append(Issue(table.filename, 1, "", "Excel dosyasında çalışma sayfası yok."))
            return
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        header_cells = next(rows, ())
        headers = [cell.value for cell in header_cells]
        while headers and headers[-1] is None:
            headers.pop()
        if not _headers(table, headers):
            return
        for number, cells in enumerate(rows, start=2):
            values = [cell.value for cell in cells]
            while len(values) > len(headers) and values[-1] is None:
                values.pop()
            values += [None] * max(0, len(headers) - len(values))
            for index, cell in enumerate(cells):
                if cell.data_type in {"f", "e"}:
                    field = headers[index] if index < len(headers) else ""
                    table.issues.append(Issue(table.filename, number, field, "Excel formülleri ve hata hücreleri desteklenmez; sabit değer kullanın."))
            if not _append_row(table, number, values):
                break
    except (BadZipFile, OSError, ValueError, KeyError, TypeError, EOFError):
        table.issues.append(Issue(table.filename, None, "", "XLSX okunamadı. Geçerli, şifresiz bir Excel dosyası kullanın."))
    except Exception:
        # XML ayrıştırıcıları farklı hata türleri üretir; kullanıcıya dosya hatası göster.
        table.issues.append(Issue(table.filename, None, "", "XLSX içeriği bozuk veya desteklenmiyor; dosyayı Excel'den yeniden kaydedin."))
    finally:
        if workbook is not None:
            workbook.close()
